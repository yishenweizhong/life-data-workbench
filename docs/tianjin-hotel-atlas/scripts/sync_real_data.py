#!/usr/bin/env python3
import argparse
import csv
import json
import re
from html.parser import HTMLParser
from pathlib import Path

from openpyxl import load_workbook


class StoreRowParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.depth = 0
        self.in_stores = False
        self.stores_depth = None
        self.rows = []

    def handle_starttag(self, tag, attrs):
        self.depth += 1
        attributes = dict(attrs)
        if tag == "section" and attributes.get("id") == "stores":
            self.in_stores = True
            self.stores_depth = self.depth
        elif self.in_stores and tag == "tr" and attributes.get("data-store"):
            self.rows.append(attributes)

    def handle_endtag(self, tag):
        if (
            self.in_stores
            and tag == "section"
            and self.depth == self.stores_depth
        ):
            self.in_stores = False
            self.stores_depth = None
        self.depth = max(0, self.depth - 1)


def number(value, default=0.0):
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def integer(value):
    return int(round(number(value)))


def load_store_rows(workbench_path):
    html = Path(workbench_path).read_text(encoding="utf-8")
    parser = StoreRowParser()
    parser.feed(html)
    rows = {}
    for row in parser.rows:
        rows[row["data-store"]] = row

    metadata_source = html.replace('\\"', '"')
    period_match = re.search(
        r'"current_start":\s*"([^"]+)".*?"current_end":\s*"([^"]+)"',
        metadata_source,
        re.DOTALL,
    )
    if not period_match:
        raise ValueError("current workbench period was not found")
    return rows, period_match.groups()


def load_score_rows(score_path):
    workbook = load_workbook(score_path, read_only=True, data_only=True)
    worksheet = workbook["门店经营分"]
    rows = worksheet.iter_rows(values_only=True)
    headers = next(rows)
    indexed = {}
    for values in rows:
        record = dict(zip(headers, values))
        poi_id = str(record.get("POI ID") or "").strip()
        if poi_id:
            indexed[poi_id] = record
    as_of = str(workbook["使用说明"]["B2"].value)
    workbook.close()
    return indexed, as_of


def load_registry(registry_path):
    payload = json.loads(Path(registry_path).read_text(encoding="utf-8"))
    return payload["hotels"] if isinstance(payload, dict) else payload


def build_dataset(workbench_path, score_path, registry_path):
    store_rows, period = load_store_rows(workbench_path)
    score_rows, as_of = load_score_rows(score_path)
    registry = load_registry(registry_path)
    hotels = []
    store_matches = 0
    score_matches = 0

    for source in registry:
        store = store_rows.get(source["name"])
        score = score_rows.get(str(source["poiId"]))
        if store:
            store_matches += 1
        if score:
            score_matches += 1
        if not store or not score:
            missing = []
            if not store:
                missing.append("store performance")
            if not score:
                missing.append("operation score")
            raise ValueError(
                f"{source['name']} is missing {' and '.join(missing)}"
            )

        hotels.append(
            {
                "id": source.get("id") or f"poi-{source['poiId']}",
                "poiId": str(source["poiId"]),
                "name": source["name"],
                "brand": source.get("brand") or "未标注",
                "district": store.get("data-area") or source.get("district"),
                "city": store.get("data-city") or "天津市",
                "address": source.get("address") or "",
                "lng": number(source.get("lng"), None),
                "lat": number(source.get("lat"), None),
                "paymentNights": integer(store.get("data-payment-nights")),
                "redeemNights": integer(store.get("data-redeem-nights")),
                "paymentGmv": round(
                    number(store.get("data-payment-current")), 2
                ),
                "redeemGmv": round(
                    number(store.get("data-redeem-current")), 2
                ),
                "paymentGrowth": number(store.get("data-payment-growth")),
                "redeemGrowth": number(store.get("data-redeem-growth")),
                "operationScore": number(score.get("经营分")),
                "scoreDelta": number(score.get("近7天变化")),
                "certification": str(score.get("金冠认证") or "未认证"),
                "scoreGaps": [
                    item.strip()
                    for item in str(score.get("未满分维度") or "").split(",")
                    if item.strip()
                ],
                "provider": store.get("data-provider") or "未绑定服务商",
                "owner": str(score.get("区域户名称") or ""),
                "updatedAt": f"{as_of}T23:59:00+08:00",
            }
        )

    if store_matches != 169 or score_matches != 169 or len(hotels) != 169:
        raise ValueError(
            "expected 169 verified Tianjin matches, got "
            f"{store_matches} store and {score_matches} score matches"
        )

    return {
        "meta": {
            "name": "津河省区天津门店真实数据",
            "source": f"生意经真实经营数据 · {as_of}",
            "sourceUrl": (
                "https://yishenweizhong.github.io/life-data-workbench/"
                "workbench_jinhe.html#stores"
            ),
            "asOf": as_of,
            "generatedAt": f"{as_of}T23:59:00+08:00",
            "period": f"{period[0]} 至 {period[1]}",
            "count": len(hotels),
            "storeMatches": store_matches,
            "scoreMatches": score_matches,
            "scope": "天津市 169 家已完成 POI 坐标匹配的真实门店",
        },
        "hotels": hotels,
    }


def write_outputs(result, app_dir):
    app_dir = Path(app_dir)
    json_path = app_dir / "data" / "tianjin-hotels.json"
    csv_path = app_dir / "data" / "tianjin-hotels-source.csv"
    js_path = app_dir / "real-data.js"

    json_text = json.dumps(
        result, ensure_ascii=False, indent=2, sort_keys=False
    )
    json_path.write_text(json_text + "\n", encoding="utf-8")

    js_text = (
        "var HOTEL_DATASET = "
        + json.dumps(
            result["meta"], ensure_ascii=False, indent=2, sort_keys=False
        )
        + ";\n\nvar HOTEL_SEED = "
        + json.dumps(
            result["hotels"], ensure_ascii=False, indent=2, sort_keys=False
        )
        + ";\n"
    )
    js_path.write_text(js_text, encoding="utf-8")

    columns = [
        ("门店名称", "name"),
        ("品牌", "brand"),
        ("二级片区", "district"),
        ("地址", "address"),
        ("经度", "lng"),
        ("纬度", "lat"),
        ("支付间夜", "paymentNights"),
        ("核销间夜", "redeemNights"),
        ("近7日支付GMV", "paymentGmv"),
        ("近7日核销GMV", "redeemGmv"),
        ("支付增长率", "paymentGrowth"),
        ("核销增长率", "redeemGrowth"),
        ("经营分", "operationScore"),
        ("近7天变化", "scoreDelta"),
        ("金冠认证", "certification"),
        ("POI ID", "poiId"),
        ("服务商", "provider"),
        ("区域户名称", "owner"),
    ]
    with csv_path.open("w", encoding="utf-8-sig", newline="") as output:
        writer = csv.writer(output, lineterminator="\n")
        writer.writerow([label for label, _ in columns])
        for hotel in result["hotels"]:
            writer.writerow([hotel.get(key, "") for _, key in columns])


def main():
    script_dir = Path(__file__).resolve().parent
    app_dir = script_dir.parent
    docs_dir = app_dir.parent
    parser = argparse.ArgumentParser(
        description="Build the verified Tianjin real-store snapshot."
    )
    parser.add_argument(
        "--workbench",
        type=Path,
        default=docs_dir / "workbench_jinhe.html",
    )
    parser.add_argument(
        "--scores",
        type=Path,
        default=docs_dir
        / "downloads"
        / "经营分_津河省区_2026-10-20.xlsx",
    )
    parser.add_argument(
        "--registry",
        type=Path,
        default=app_dir / "data" / "tianjin-hotels.json",
    )
    args = parser.parse_args()
    result = build_dataset(args.workbench, args.scores, args.registry)
    write_outputs(result, app_dir)
    print(
        f"Wrote {result['meta']['count']} verified stores "
        f"for {result['meta']['asOf']}"
    )


if __name__ == "__main__":
    main()
