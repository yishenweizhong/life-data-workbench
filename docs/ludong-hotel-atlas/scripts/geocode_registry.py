#!/usr/bin/env python3
import argparse
import json
import re
import threading
import unicodedata
from difflib import SequenceMatcher
from functools import partial
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

from sync_real_data import load_score_rows, load_store_rows


DEFAULT_ORG_PATH = Path(
    "/Users/bytedance/.trae-cn/attachments/"
    "fb6cb173-dd31-49d7-9405-706a92c935bf_"
    "c79f7fa4-db40-4a8a-ba9e-c0604102b106_1018-2026-10-18 14-51-32.xlsx"
)
DEFAULT_LEDGER_PATH = Path(
    "/Users/bytedance/.trae-cn/attachments/"
    "3cb87ed5-f72d-4046-8fe3-20c9b62d570f_"
    "67fd5973-47d3-4289-82a3-f633e12befc3_华北区抖音数据台账.xlsx"
)


def normalize(value):
    text = unicodedata.normalize("NFKC", str(value or "")).lower()
    return re.sub(r"[^a-z0-9\u4e00-\u9fff]", "", text)


def brand_from_name(name):
    clean = re.split(r"[（(·]", name, maxsplit=1)[0]
    return clean.replace("酒店", "").strip() or "未标注"


def load_organizations(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    headers = next(rows)
    records = []
    for values in rows:
        record = dict(zip(headers, values))
        if record.get("门店户商家名称") in (None, "总计"):
            continue
        provider = str(
            record.get("门店户平台代运营服务商名称")
            or "未绑定服务商"
        ).strip()
        if provider.upper() in ("NULL", "N/A", "#N/A"):
            provider = "未绑定服务商"
        records.append(
            {
                "name": str(record.get("门店户商家名称") or "").strip(),
                "city": str(record.get("POI所在城市") or "").strip(),
                "provider": provider,
                "poiId": str(record.get("POI_ID") or "").strip(),
            }
        )
    workbook.close()
    return records


def load_legacy_names(path, poi_ids):
    if not path.exists():
        return {}
    workbook = load_workbook(path, read_only=True, data_only=True)
    candidates = {poi_id: [] for poi_id in poi_ids}
    for worksheet in workbook.worksheets:
        for values in worksheet.iter_rows(values_only=True):
            row = [str(value).strip() if value is not None else "" for value in values]
            if not row or row[0] not in candidates:
                continue
            candidates[row[0]].extend(row[1:4])
    workbook.close()

    def quality(value):
        chinese = len(re.findall(r"[\u4e00-\u9fff]", value))
        hotel_bonus = 50 if "酒店" in value else 0
        code_penalty = -100 if re.fullmatch(r"[A-Z0-9]+", value) else 0
        return hotel_bonus + chinese * 3 + min(len(value), 50) + code_penalty

    return {
        poi_id: max((value for value in values if value), key=quality)
        for poi_id, values in candidates.items()
        if any(values)
    }


def infer_city(name, district):
    city_tokens = {
        "青岛": "青岛市",
        "烟台": "烟台市",
        "威海": "威海市",
        "日照": "日照市",
        "临沂": "临沂市",
        "枣庄": "枣庄市",
        "淄博": "淄博市",
        "潍坊": "潍坊市",
        "滕州": "枣庄市",
        "蓬莱": "烟台市",
        "龙口": "烟台市",
        "招远": "烟台市",
        "海阳": "烟台市",
        "莱州": "烟台市",
        "乳山": "威海市",
        "荣成": "威海市",
        "诸城": "潍坊市",
        "寿光": "潍坊市",
        "青州": "潍坊市",
        "高密": "潍坊市",
        "安丘": "潍坊市",
        "临朐": "潍坊市",
        "青州市": "潍坊市",
        "青州": "潍坊市",
        "莒县": "日照市",
        "沂水": "临沂市",
        "郯城": "临沂市",
        "临沭": "临沂市",
        "蒙阴": "临沂市",
        "费县": "临沂市",
        "台儿庄": "枣庄市",
        "临淄": "淄博市",
    }
    for token, city in city_tokens.items():
        if token in name:
            return city
    district_defaults = {
        "临沂片区": "临沂市",
        "威烟片区": "威海市",
        "日照片区": "日照市",
        "沂枣片区": "枣庄市",
        "淄博片区": "淄博市",
        "潍坊片区": "潍坊市",
        "烟台片区": "烟台市",
        "青岛前海片区": "青岛市",
        "青岛北片区": "青岛市",
        "青岛西片区": "青岛市",
    }
    return district_defaults.get(district, "")


def assign_exception_districts(records):
    city_district = {
        "临沂市": "临沂片区",
        "威海市": "威烟片区",
        "日照市": "日照片区",
        "枣庄市": "沂枣片区",
        "淄博市": "淄博片区",
        "潍坊市": "潍坊片区",
        "烟台市": "烟台片区",
    }
    named_qingdao = [
        record for record in records[:329]
        if record.get("city") == "青岛市"
        and record.get("lng") is not None
        and record.get("lat") is not None
    ]
    pending_districts = iter(["沂枣片区", "青岛北片区", "青岛西片区"])
    for record in records[329:]:
        if record["matchStatus"] == "identity-pending":
            record["district"] = next(pending_districts)
            record["city"] = (
                "枣庄市" if record["district"] == "沂枣片区" else "青岛市"
            )
            continue
        if record["city"] in city_district:
            record["district"] = city_district[record["city"]]
            continue
        if (
            record["city"] == "青岛市"
            and record.get("lng") is not None
            and record.get("lat") is not None
            and named_qingdao
        ):
            nearest = min(
                named_qingdao,
                key=lambda item: (
                    (item["lng"] - record["lng"]) ** 2
                    + (item["lat"] - record["lat"]) ** 2
                ),
            )
            record["district"] = nearest["district"]
        elif record["city"] == "青岛市":
            record["district"] = "青岛北片区"
    return records


def match_named_rows(store_rows, score_rows, organizations):
    by_name = {}
    for record in organizations:
        by_name.setdefault(normalize(record["name"]), []).append(record)
    score_ids = set(score_rows)
    available = [record for record in organizations if record["poiId"] in score_ids]
    matches = []
    used_ids = set()

    for row_index, store in enumerate(store_rows[:329]):
        candidates = [
            item for item in by_name.get(normalize(store["data-store"]), [])
            if item["poiId"] not in used_ids and item["poiId"] in score_ids
        ]
        if not candidates:
            same_city = [
                item for item in available
                if item["poiId"] not in used_ids
                and (
                    not store.get("data-city")
                    or item["city"] == store.get("data-city")
                )
            ]
            ranked = sorted(
                same_city,
                key=lambda item: SequenceMatcher(
                    None,
                    normalize(store["data-store"]),
                    normalize(item["name"]),
                ).ratio(),
                reverse=True,
            )
            candidates = ranked[:1]
        if len(candidates) != 1:
            raise ValueError(f"organization match failed for {store['data-store']}")
        match = dict(candidates[0])
        similarity = SequenceMatcher(
            None, normalize(store["data-store"]), normalize(match["name"])
        ).ratio()
        if similarity < 0.70:
            raise ValueError(
                f"low-confidence organization match for {store['data-store']}: "
                f"{match['name']} ({similarity:.2f})"
            )
        used_ids.add(match["poiId"])
        match.update(
            {
                "sourceRowIndex": row_index,
                "sourceName": store["data-store"],
                "district": store["data-area"],
                "city": match["city"] or store.get("data-city") or "",
                "brand": brand_from_name(match["name"]),
                "matchStatus": "organization-exact"
                if similarity == 1
                else "organization-fuzzy",
            }
        )
        matches.append(match)
    if len(matches) != 329:
        raise ValueError(f"expected 329 named matches, got {len(matches)}")
    return matches


def build_registry(workbench_path, score_path, organization_path, ledger_path):
    store_rows, _ = load_store_rows(workbench_path)
    score_rows, _, exception_ids = load_score_rows(score_path)
    organizations = load_organizations(organization_path)
    records = match_named_rows(store_rows, score_rows, organizations)
    legacy_names = load_legacy_names(ledger_path, exception_ids)

    for offset, poi_id in enumerate(exception_ids):
        source_index = 329 + offset
        source = store_rows[source_index]
        recovered = legacy_names.get(poi_id)
        name = recovered or f"待补门店（POI {poi_id}）"
        city = {
            "7263417326605715514": "潍坊市",
        }.get(poi_id, infer_city(name, ""))
        records.append(
            {
                "poiId": poi_id,
                "name": name,
                "sourceName": source["data-store"],
                "sourceRowIndex": source_index,
                "district": "",
                "city": city,
                "provider": "未绑定服务商",
                "brand": brand_from_name(name) if recovered else "待补充",
                "matchStatus": "legacy-name" if recovered else "identity-pending",
            }
        )
    if len(records) != 387:
        raise ValueError(f"expected 387 registry records, got {len(records)}")
    return records


def choose_poi(record, pois):
    expected = normalize(record["name"])
    ranked = []
    for poi in pois or []:
        similarity = SequenceMatcher(
            None, expected, normalize(poi.get("name"))
        ).ratio()
        location = poi.get("location") or {}
        if isinstance(location, list):
            lng, lat = location
        else:
            lng, lat = location.get("lng"), location.get("lat")
        if lng is None or lat is None:
            continue
        ranked.append((similarity, poi, float(lng), float(lat)))
    if not ranked:
        return None
    ranked.sort(key=lambda item: item[0], reverse=True)
    similarity, poi, lng, lat = ranked[0]
    if similarity < 0.38:
        return None
    return {
        "amapName": poi.get("name") or "",
        "address": poi.get("address") or "",
        "lng": lng,
        "lat": lat,
        "geocodeSimilarity": round(similarity, 4),
        "geocodeStatus": "matched" if similarity >= 0.55 else "review",
    }


def geocode(records, docs_dir, page_url=None):
    server = None
    if page_url:
        url = page_url
    else:
        handler = partial(SimpleHTTPRequestHandler, directory=str(docs_dir))
        server = ThreadingHTTPServer(("127.0.0.1", 4176), handler)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        url = "http://127.0.0.1:4176/tianjin-hotel-atlas/"
    try:
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(headless=True)
            page = browser.new_page()
            amap_ready = False
            for _ in range(3):
                page.goto(url, wait_until="domcontentloaded", timeout=30000)
                page.wait_for_timeout(5000)
                amap_ready = page.evaluate(
                    "typeof window.AMap === 'object' && "
                    "typeof window.AMap.plugin === 'function'"
                )
                if amap_ready:
                    break
            if not amap_ready:
                raise RuntimeError("AMap JavaScript API did not load after 3 attempts")
            page.evaluate(
                """() => new Promise((resolve) => {
                  AMap.plugin("AMap.PlaceSearch", resolve);
                })"""
            )
            for index, record in enumerate(records, start=1):
                if record.get("lng") is not None and record.get("lat") is not None:
                    continue
                if record["matchStatus"] == "identity-pending":
                    record.update(
                        {
                            "address": "",
                            "lng": None,
                            "lat": None,
                            "geocodeStatus": "identity-pending",
                        }
                    )
                    continue
                payload = page.evaluate(
                    """async ({name, city}) => Promise.race([
                      new Promise((resolve) => {
                        const search = new AMap.PlaceSearch({
                          city: city || "",
                          citylimit: Boolean(city),
                          pageSize: 10,
                          pageIndex: 1,
                          extensions: "all"
                        });
                        search.search(name, (status, result) => {
                          const list = result && result.poiList;
                          resolve({
                            status,
                            info: result && result.info,
                            pois: list ? list.pois.map((poi) => ({
                              name: poi.name,
                              address: poi.address,
                              location: poi.location
                                ? [poi.location.lng, poi.location.lat]
                                : null
                            })) : []
                          });
                        });
                      }),
                      new Promise((resolve) => setTimeout(
                        () => resolve({status: "timeout", pois: []}), 12000
                      ))
                    ])""",
                    {"name": record["name"], "city": record["city"]},
                )
                selected = choose_poi(record, payload.get("pois"))
                if not selected and payload.get("status") in ("error", "timeout"):
                    page.wait_for_timeout(900)
                    payload = page.evaluate(
                        """async ({name, city}) => new Promise((resolve) => {
                          const search = new AMap.PlaceSearch({
                            city: city || "",
                            citylimit: Boolean(city),
                            pageSize: 10,
                            pageIndex: 1,
                            extensions: "all"
                          });
                          search.search(name, (status, result) => {
                            const list = result && result.poiList;
                            resolve({
                              status,
                              pois: list ? list.pois.map((poi) => ({
                                name: poi.name,
                                address: poi.address,
                                location: poi.location
                                  ? [poi.location.lng, poi.location.lat]
                                  : null
                              })) : []
                            });
                          });
                        })""",
                        {"name": record["name"], "city": record["city"]},
                    )
                    selected = choose_poi(record, payload.get("pois"))
                if selected:
                    record.update(selected)
                else:
                    record.update(
                        {
                            "address": "",
                            "lng": None,
                            "lat": None,
                            "geocodeStatus": payload.get("status") or "unmatched",
                        }
                    )
                if index % 25 == 0:
                    print(f"geocoded {index}/{len(records)}")
                page.wait_for_timeout(280)
            browser.close()
    finally:
        if server:
            server.shutdown()
            server.server_close()
    return assign_exception_districts(records)


def main():
    script_dir = Path(__file__).resolve().parent
    app_dir = script_dir.parent
    docs_dir = app_dir.parent
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--workbench",
        type=Path,
        default=app_dir / "data" / "ludong-workbench-2026-10-20.json",
    )
    parser.add_argument(
        "--scores",
        type=Path,
        default=docs_dir / "downloads" / "经营分_鲁东省区_2026-10-20.xlsx",
    )
    parser.add_argument("--organizations", type=Path, default=DEFAULT_ORG_PATH)
    parser.add_argument("--ledger", type=Path, default=DEFAULT_LEDGER_PATH)
    parser.add_argument(
        "--output",
        type=Path,
        default=app_dir / "data" / "ludong-hotel-registry.json",
    )
    parser.add_argument("--page-url")
    parser.add_argument("--skip-geocode", action="store_true")
    args = parser.parse_args()

    records = build_registry(
        args.workbench, args.scores, args.organizations, args.ledger
    )
    if args.output.exists():
        previous = json.loads(args.output.read_text(encoding="utf-8"))
        previous_by_id = {
            record["poiId"]: record for record in previous.get("hotels", [])
        }
        for record in records:
            old = previous_by_id.get(record["poiId"], {})
            if (
                old.get("lng") is not None
                and old.get("lat") is not None
                and old.get("city") == record.get("city")
            ):
                for key in (
                    "amapName",
                    "address",
                    "lng",
                    "lat",
                    "geocodeSimilarity",
                    "geocodeStatus",
                ):
                    record[key] = old.get(key)
    if not args.skip_geocode:
        records = geocode(records, docs_dir, args.page_url)
    payload = {
        "meta": {
            "count": len(records),
            "located": sum(record.get("lng") is not None for record in records),
        },
        "hotels": records,
    }
    args.output.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(
        f"Wrote {payload['meta']['count']} registry records, "
        f"{payload['meta']['located']} located"
    )


if __name__ == "__main__":
    main()
