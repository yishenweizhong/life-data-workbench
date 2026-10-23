#!/usr/bin/env python3
import argparse
import csv
import json
import re
from html.parser import HTMLParser
from pathlib import Path

from openpyxl import load_workbook


DEFAULT_ORG_PATH = Path(
    "/Users/bytedance/.trae-cn/attachments/"
    "fb6cb173-dd31-49d7-9405-706a92c935bf_"
    "c79f7fa4-db40-4a8a-ba9e-c0604102b106_1018-2026-10-18 14-51-32.xlsx"
)


class StoreRowParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.rows = []

    def handle_starttag(self, tag, attrs):
        attributes = dict(attrs)
        if tag == "tr" and attributes.get("data-store"):
            self.rows.append(attributes)


def number(value, default=0.0):
    if value is None or value == "":
        return default
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def integer(value):
    return int(round(number(value)))


def load_store_rows(workbench_path):
    if Path(workbench_path).suffix == ".json":
        payload = json.loads(Path(workbench_path).read_text(encoding="utf-8"))
        rows = payload["rows"]
        period = payload["meta"]["period"]
        if len(rows) != 387:
            raise ValueError(f"expected 387 workbench rows, got {len(rows)}")
        return rows, period
    html = Path(workbench_path).read_text(encoding="utf-8")
    parser = StoreRowParser()
    parser.feed(html)
    metadata_source = html.replace('\\"', '"')
    period_match = re.search(
        r'"current_start":\s*"([^"]+)".*?"current_end":\s*"([^"]+)"',
        metadata_source,
        re.DOTALL,
    )
    if len(parser.rows) != 387:
        raise ValueError(f"expected 387 workbench rows, got {len(parser.rows)}")
    if not period_match:
        raise ValueError("current workbench period was not found")
    return parser.rows, period_match.groups()


def load_score_rows(score_path):
    workbook = load_workbook(score_path, read_only=True, data_only=True)
    worksheet = workbook["门店经营分"]
    rows = worksheet.iter_rows(values_only=True)
    headers = next(rows)
    indexed = {}
    for values in rows:
        record = dict(zip(headers, values))
        poi_id = str(record.get("POI ID") or "").strip()
        if poi_id:
            indexed[poi_id] = record

    exception_rows = workbook["匹配异常"].iter_rows(values_only=True)
    exception_headers = next(exception_rows)
    exception_ids = []
    for values in exception_rows:
        record = dict(zip(exception_headers, values))
        poi_id = str(record.get("POI ID") or "").strip()
        if poi_id:
            exception_ids.append(poi_id)
    as_of = str(workbook["使用说明"]["B2"].value)
    workbook.close()
    return indexed, as_of, exception_ids


def load_registry(registry_path):
    payload = json.loads(Path(registry_path).read_text(encoding="utf-8"))
    return payload["hotels"] if isinstance(payload, dict) else payload


def load_organization_ids(organization_path):
    workbook = load_workbook(organization_path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    headers = next(rows)
    ids = set()
    for values in rows:
        record = dict(zip(headers, values))
        poi_id = str(record.get("POI_ID") or "").strip()
        if poi_id:
            ids.add(poi_id)
    workbook.close()
    return ids


def build_dataset(
    workbench_path,
    score_path,
    organization_path,
    registry_path,
):
    store_rows, period = load_store_rows(workbench_path)
    score_rows, as_of, exception_ids = load_score_rows(score_path)
    organization_ids = load_organization_ids(organization_path)
    registry = load_registry(registry_path)
    if len(registry) != 387:
        raise ValueError(f"expected 387 registry records, got {len(registry)}")

    hotels = []
    score_matches = 0
    score_missing = 0
    for source in registry:
        row_index = int(source["sourceRowIndex"])
        store = store_rows[row_index]
        poi_id = str(source["poiId"])
        score = score_rows.get(poi_id)
        if score:
            score_matches += 1
        else:
            score_missing += 1
            if poi_id not in exception_ids:
                raise ValueError(f"unexpected missing score for {poi_id}")

        name = str(source.get("name") or "").strip()
        if not name or name == "#N/A":
            raise ValueError(f"invalid registry name for {poi_id}")
        city = str(source.get("city") or store.get("data-city") or "").strip()
        hotels.append(
            {
                "id": f"poi-{poi_id}",
                "poiId": poi_id,
                "name": name,
                "brand": source.get("brand") or "未标注",
                "district": source.get("district") or store.get("data-area"),
                "city": city or "城市待确认",
                "address": source.get("address") or "",
                "lng": number(source.get("lng"), None),
                "lat": number(source.get("lat"), None),
                "coordinateStatus": source.get("geocodeStatus") or "unmatched",
                "paymentNights": integer(store.get("data-payment-nights")),
                "redeemNights": integer(store.get("data-redeem-nights")),
                "paymentGmv": round(number(store.get("data-payment-current")), 2),
                "redeemGmv": round(number(store.get("data-redeem-current")), 2),
                "paymentGrowth": number(store.get("data-payment-growth")),
                "redeemGrowth": number(store.get("data-redeem-growth")),
                "operationScore": (
                    number(score.get("经营分"), None) if score else None
                ),
                "scoreDelta": (
                    number(score.get("近7天变化"), None) if score else None
                ),
                "certification": (
                    str(score.get("金冠认证") or "未认证") if score else "暂缺"
                ),
                "scoreGaps": (
                    [
                        item.strip()
                        for item in str(score.get("未满分维度") or "").split(",")
                        if item.strip()
                    ]
                    if score
                    else []
                ),
                "provider": source.get("provider")
                or store.get("data-provider")
                or "未绑定服务商",
                "owner": str(score.get("区域户名称") or "") if score else "",
                "updatedAt": f"{as_of}T23:59:00+08:00",
            }
        )

    if score_matches != 329 or score_missing != 58:
        raise ValueError(
            f"expected 329 scored and 58 missing, got "
            f"{score_matches} and {score_missing}"
        )
    if sum(hotel["poiId"] in organization_ids for hotel in hotels) != 329:
        raise ValueError("organization ledger should match all 329 scored stores")

    return {
        "meta": {
            "name": "鲁东省区酒店真实经营数据",
            "source": f"生意经真实经营数据 · {as_of}",
            "sourceUrl": (
                "https://yishenweizhong.github.io/life-data-workbench/"
                "workbench_ludong.html#stores"
            ),
            "asOf": as_of,
            "generatedAt": f"{as_of}T23:59:00+08:00",
            "period": f"{period[0]} 至 {period[1]}",
            "count": len(hotels),
            "scoreMatches": score_matches,
            "scoreMissing": score_missing,
            "located": sum(hotel["lng"] is not None for hotel in hotels),
            "scope": "鲁东省区 387 家真实组织门店",
        },
        "hotels": hotels,
    }


def write_outputs(result, app_dir):
    app_dir = Path(app_dir)
    json_path = app_dir / "data" / "ludong-hotels.json"
    csv_path = app_dir / "data" / "ludong-hotels-source.csv"
    js_path = app_dir / "real-data.js"
    json_text = json.dumps(result, ensure_ascii=False, indent=2)
    json_path.write_text(json_text + "\n", encoding="utf-8")
    js_path.write_text(
        "var HOTEL_DATASET = "
        + json.dumps(result["meta"], ensure_ascii=False, indent=2)
        + ";\n\nvar HOTEL_SEED = "
        + json.dumps(result["hotels"], ensure_ascii=False, indent=2)
        + ";\n",
        encoding="utf-8",
    )

    columns = [
        ("门店名称", "name"),
        ("品牌", "brand"),
        ("二级片区", "district"),
        ("城市", "city"),
        ("地址", "address"),
        ("经度", "lng"),
        ("纬度", "lat"),
        ("支付间夜", "paymentNights"),
        ("核销间夜", "redeemNights"),
        ("近7日支付GMV", "paymentGmv"),
        ("近7日核销GMV", "redeemGmv"),
        ("支付增长率", "paymentGrowth"),
        ("核销增长率", "redeemGrowth"),
        ("经营分", "operationScore"),
        ("近7天变化", "scoreDelta"),
        ("金冠认证", "certification"),
        ("POI ID", "poiId"),
        ("服务商", "provider"),
        ("区域户名称", "owner"),
    ]
    with csv_path.open("w", encoding="utf-8-sig", newline="") as output:
        writer = csv.writer(output, lineterminator="\n")
        writer.writerow([label for label, _ in columns])
        for hotel in result["hotels"]:
            writer.writerow(
                [
                    "" if hotel.get(key) is None else hotel.get(key, "")
                    for _, key in columns
                ]
            )


def main():
    script_dir = Path(__file__).resolve().parent
    app_dir = script_dir.parent
    docs_dir = app_dir.parent
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--workbench",
        type=Path,
        default=app_dir / "data" / "ludong-workbench-2026-10-20.json",
    )
    parser.add_argument(
        "--scores",
        type=Path,
        default=docs_dir / "downloads" / "经营分_鲁东省区_2026-10-20.xlsx",
    )
    parser.add_argument(
        "--organizations", type=Path, default=DEFAULT_ORG_PATH
    )
    parser.add_argument(
        "--registry",
        type=Path,
        default=app_dir / "data" / "ludong-hotel-registry.json",
    )
    args = parser.parse_args()
    result = build_dataset(
        args.workbench, args.scores, args.organizations, args.registry
    )
    write_outputs(result, app_dir)
    print(
        f"Wrote {result['meta']['count']} stores, "
        f"{result['meta']['scoreMatches']} scores, "
        f"{result['meta']['scoreMissing']} missing scores"
    )


if __name__ == "__main__":
    main()
